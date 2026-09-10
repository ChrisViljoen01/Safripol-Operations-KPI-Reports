"""
Builds "SAF Ops Tracking - TAC IMOLA.xlsx": the operational capture workbook for the
PTA outbound drawdown.

Design principles
-----------------
1. Capture once, calculate everywhere. Anything derivable is a formula, never typed.
2. Real datetimes for Start/End so cross-midnight nightshift decants compute correctly.
   The old tracker stored time-only values, which broke every duration that spanned 00:00.
3. Delay time is a NUMBER (minutes) with a controlled category, not free text.
   The old "Time Loss/Saved" column held "Clean"/"Time Exceeded" and could not be summed.
4. Turnaround (transit / Safripol dwell / return) is NOT captured by hand. It is derived
   from the GPS dwell feed and the stock report, so the yard team only captures decanting.
5. Every list field is validated, so the Power BI / HTML report never has to clean text.

Run:  python tools/build_tracker.py --out "<path>\\SAF Ops Tracking - TAC IMOLA.xlsx"
"""

from __future__ import annotations

import argparse
import os
from datetime import date

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

# ---------------------------------------------------------------------------
# Shipment constants
# ---------------------------------------------------------------------------
SHIPMENT = "TAC IMOLA"
TARGET_TOTAL_MT = 20390.40          # 16,992 bags x 1.2 MT  (admin basis)
TARGET_LOAD_KG = 21600              # nominal isotainer payload
ALLOWED_DECANT_HRS = 1.75           # 1h45 standard
ALLOWED_INTERVAL_HRS = 1.00         # target gap between decants
SITE_TARGET_HRS = 2.75              # dwell target both sites
TRANSIT_TARGET_HRS = 0.75
ROWS = 4000                         # pre-formulated capture rows

INK = "0B1220"
ACCENT = "1E4FD8"
HEAD_FILL = PatternFill("solid", fgColor="12224A")
SUB_FILL = PatternFill("solid", fgColor="EAF0FF")
LOCK_FILL = PatternFill("solid", fgColor="F2F4F8")
WARN_FILL = PatternFill("solid", fgColor="FFE0E0")
GOOD_FILL = PatternFill("solid", fgColor="DFF5E5")
THIN = Side(style="thin", color="C7D0E0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

DELAY_CATEGORIES = [
    "None",
    "Bad / Damaged Bags",
    "Breakdown - Forklift",
    "Breakdown - Straddle Carrier",
    "Breakdown - Other Equipment",
    "Diesel / Refuelling",
    "Isotainer Not Available",
    "Isotainer Inspection Failure",
    "Truck / Driver Not Available",
    "Safripol Site Delay",
    "Safripol Offload Queue",
    "Weather - Rain",
    "Weather - Wind",
    "Power Outage / Loadshedding",
    "WIFI / System Downtime",
    "Weighbridge Delay",
    "Labour Shortage",
    "Shift Handover",
    "Meal / Statutory Break",
    "Housekeeping / Cleaning",
    "Safety Stoppage",
    "Waiting on Instruction",
    "Other",
]

DELAY_OWNERS = ["Connect", "Safripol", "Transporter", "Third Party", "External / Weather"]
SHIFTS = ["Dayshift", "Nightshift"]
TEAMS = ["A", "B", "C"]
ISOS = ["ISO-01", "ISO-02", "ISO-03", "ISO-04", "ISO-05"]
YESNO = ["Yes", "No"]


def _list_dv(items: list[str], prompt: str) -> DataValidation:
    dv = DataValidation(
        type="list",
        formula1='"' + ",".join(items) + '"',
        allow_blank=True,
        showDropDown=False,
    )
    dv.error = "Select a value from the list."
    dv.errorTitle = "Invalid entry"
    dv.prompt = prompt
    dv.promptTitle = "Select"
    return dv


def _style_header(ws, row: int = 1, upto: int | None = None) -> None:
    upto = upto or ws.max_column
    for c in range(1, upto + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEAD_FILL
        cell.font = Font(bold=True, color="FFFFFF", size=10)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.row_dimensions[row].height = 34
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def _widths(ws, widths: dict[str, int]) -> None:
    for col, w in widths.items():
        ws.column_dimensions[col].width = w


# ---------------------------------------------------------------------------
# 1. Lookups / control sheet
# ---------------------------------------------------------------------------
def build_lookups(wb: Workbook):
    ws = wb.create_sheet("Lookups")
    ws["A1"] = "CONTROL PARAMETERS"
    ws["A1"].font = Font(bold=True, size=12, color=ACCENT)

    params = [
        ("Shipment", SHIPMENT, "Vessel / shipment reference"),
        ("Target Total MT", TARGET_TOTAL_MT, "16,992 bags x 1.2 MT - Safripol delivery basis"),
        ("Target Load KG", TARGET_LOAD_KG, "Nominal isotainer payload"),
        ("Allowed Decant Hours", ALLOWED_DECANT_HRS, "Standard 1h45 per isotainer"),
        ("Allowed Interval Hours", ALLOWED_INTERVAL_HRS, "Target gap between consecutive decants"),
        ("Site Dwell Target Hours", SITE_TARGET_HRS, "Connect and Safripol dwell target"),
        ("Transit Target Hours", TRANSIT_TARGET_HRS, "Connect <-> Safripol one-way target"),
        ("Bag Weight MT", 1.2, "Administrative bag conversion"),
    ]
    ws["A3"], ws["B3"], ws["C3"] = "Parameter", "Value", "Description"
    for c in ("A3", "B3", "C3"):
        ws[c].font = Font(bold=True, color="FFFFFF")
        ws[c].fill = HEAD_FILL
    for i, (k, v, d) in enumerate(params, start=4):
        ws.cell(row=i, column=1, value=k).font = Font(bold=True)
        ws.cell(row=i, column=2, value=v)
        ws.cell(row=i, column=3, value=d).font = Font(italic=True, color="667085")

    # Named ranges so formulas stay readable and one edit re-rates the whole book.
    for idx, (k, _, _) in enumerate(params, start=4):
        name = k.replace(" ", "_")
        wb.defined_names.add(__import__("openpyxl").workbook.defined_name.DefinedName(
            name, attr_text=f"Lookups!$B${idx}"))

    ws["E3"] = "Delay Categories"
    ws["E3"].font = Font(bold=True, color="FFFFFF")
    ws["E3"].fill = HEAD_FILL
    for i, v in enumerate(DELAY_CATEGORIES, start=4):
        ws.cell(row=i, column=5, value=v)

    ws["G3"] = "Delay Owner"
    ws["G3"].font = Font(bold=True, color="FFFFFF")
    ws["G3"].fill = HEAD_FILL
    for i, v in enumerate(DELAY_OWNERS, start=4):
        ws.cell(row=i, column=7, value=v)

    ws["I3"] = "Isotainers"
    ws["I3"].font = Font(bold=True, color="FFFFFF")
    ws["I3"].fill = HEAD_FILL
    for i, v in enumerate(ISOS, start=4):
        ws.cell(row=i, column=9, value=v)

    _widths(ws, {"A": 26, "B": 16, "C": 58, "E": 30, "G": 22, "I": 14})
    return ws


# ---------------------------------------------------------------------------
# 2. Decant Log  - the sheet the yard team actually fills in
# ---------------------------------------------------------------------------
DECANT_COLS = [
    # (header, key, width, kind)  kind: input | calc | note
    ("Shift Date",                  "A", 12, "input"),
    ("Shift",                       "B", 12, "input"),
    ("Team",                        "C", 8,  "input"),
    ("Heads on Shift",              "D", 9,  "input"),
    ("Isotainer",                   "E", 11, "input"),
    ("Start Date & Time",           "F", 18, "input"),
    ("End Date & Time",             "G", 18, "input"),
    ("Tare Weight (kg)",            "H", 11, "input"),
    ("Empty Weight (kg)",           "I", 11, "input"),
    ("Loaded Weight (kg)",          "J", 12, "input"),
    ("Delay Minutes",               "K", 10, "input"),
    ("Delay Category",              "L", 24, "input"),
    ("Delay Owner",                 "M", 16, "input"),
    ("Delay Detail",                "N", 34, "input"),
    ("Escalate?",                   "O", 10, "input"),
    ("Comments",                    "P", 30, "input"),
    # ---- calculated from here ----
    ("Month",                       "Q", 11, "calc"),
    ("Day",                         "R", 11, "calc"),
    ("Decant Time (hrs)",           "S", 11, "calc"),
    ("Decant Interval (hrs)",       "T", 12, "calc"),
    ("Allowed Time (hrs)",          "U", 11, "calc"),
    ("Time Variance (hrs)",         "V", 11, "calc"),
    ("Performance",                 "W", 14, "calc"),
    ("Carry Over Var (kg)",         "X", 12, "calc"),
    ("Excess Weight Var (kg)",      "Y", 13, "calc"),
    ("Net Product (kg)",            "Z", 12, "calc"),
    ("Net Product (MT)",            "AA", 12, "calc"),
    ("Productive Hours",            "AB", 11, "calc"),
    ("Shift Key",                   "AC", 20, "calc"),
]


def build_decant_log(wb: Workbook):
    ws = wb.create_sheet("Decant Log", 0)

    for header, col, width, kind in DECANT_COLS:
        cell = ws[f"{col}1"]
        cell.value = header
        ws.column_dimensions[col].width = width
    _style_header(ws, 1, upto=len(DECANT_COLS))

    # tint the calculated block so operators know not to type there
    for header, col, width, kind in DECANT_COLS:
        if kind == "calc":
            ws[f"{col}1"].fill = PatternFill("solid", fgColor="334155")

    for r in range(2, ROWS + 2):
        p = r  # row alias
        # Q Month / R Day
        ws[f"Q{p}"] = f'=IF($A{p}="","",TEXT($A{p},"mmmm"))'
        ws[f"R{p}"] = f'=IF($A{p}="","",TEXT($A{p},"dddd"))'
        # S Decant time in hours (datetime aware -> handles nightshift roll-over)
        ws[f"S{p}"] = (
            f'=IF(OR($F{p}="",$G{p}=""),"",'
            f'IF($G{p}<$F{p},"CHECK TIMES",ROUND(($G{p}-$F{p})*24,3)))'
        )
        # T Interval = start of this decant less the end of the PREVIOUS decant on the
        #   same shift. Uses the largest End time that is earlier than this Start.
        #   MAXIFS is used rather than LET so the book works on Excel 2019 and later.
        prev_end = (
            f'MAXIFS($G$2:$G${ROWS + 1},$G$2:$G${ROWS + 1},"<"&$F{p},'
            f'$AC$2:$AC${ROWS + 1},$AC{p})'
        )
        ws[f"T{p}"] = (
            f'=IF($F{p}="","",IFERROR(IF({prev_end}=0,"",'
            f'ROUND(($F{p}-{prev_end})*24,3)),""))'
        )
        # U Allowed
        ws[f"U{p}"] = f'=IF($F{p}="","",Allowed_Decant_Hours)'
        # V Variance (+ve = time lost against standard)
        ws[f"V{p}"] = f'=IF(OR($S{p}="",$S{p}="CHECK TIMES"),"",ROUND($S{p}-$U{p},3))'
        # W Performance
        ws[f"W{p}"] = (
            f'=IF($V{p}="","",IF($V{p}<=0,"Clean","Time Exceeded"))'
        )
        # X Carry over variance
        ws[f"X{p}"] = f'=IF(OR($H{p}="",$I{p}=""),"",$I{p}-$H{p})'
        # Y Excess vs nominal payload
        ws[f"Y{p}"] = f'=IF($J{p}="","",$J{p}-Target_Load_KG)'
        # Z Net product actually in the tank = loaded gross less empty tare
        ws[f"Z{p}"] = f'=IF(OR($J{p}="",$I{p}=""),"",$J{p}-$I{p})'
        ws[f"AA{p}"] = f'=IF($Z{p}="","",ROUND($Z{p}/1000,3))'
        # AB productive hours = decant time less recorded delay
        ws[f"AB{p}"] = (
            f'=IF($S{p}="","",IF($S{p}="CHECK TIMES","",'
            f'MAX(0,ROUND($S{p}-N($K{p})/60,3))))'
        )
        # AC shift key ties a row to Shift Plan
        ws[f"AC{p}"] = f'=IF($A{p}="","",TEXT($A{p},"yyyy-mm-dd")&"|"&$B{p})'

        for header, col, width, kind in DECANT_COLS:
            c = ws[f"{col}{p}"]
            c.border = BORDER
            c.font = Font(size=10)
            if kind == "calc":
                c.fill = LOCK_FILL

        ws[f"A{p}"].number_format = "dd/mm/yyyy"
        ws[f"F{p}"].number_format = "dd/mm/yyyy hh:mm"
        ws[f"G{p}"].number_format = "dd/mm/yyyy hh:mm"
        for col in ("S", "T", "U", "V", "AB"):
            ws[f"{col}{p}"].number_format = "0.00"
        for col in ("H", "I", "J", "X", "Y", "Z"):
            ws[f"{col}{p}"].number_format = "#,##0"
        ws[f"AA{p}"].number_format = "#,##0.000"

    dvs = {
        "B": _list_dv(SHIFTS, "Dayshift or Nightshift"),
        "C": _list_dv(TEAMS, "Crew team"),
        "E": _list_dv(ISOS, "Isotainer unit"),
        "M": _list_dv(DELAY_OWNERS, "Who owns the delay"),
        "O": _list_dv(YESNO, "Escalate to management?"),
    }
    for col, dv in dvs.items():
        ws.add_data_validation(dv)
        dv.add(f"{col}2:{col}{ROWS + 1}")

    dv_cat = DataValidation(
        type="list", formula1=f"=Lookups!$E$4:$E${3 + len(DELAY_CATEGORIES)}",
        allow_blank=True, showDropDown=False)
    dv_cat.error = "Pick a delay category from the list."
    ws.add_data_validation(dv_cat)
    dv_cat.add(f"L2:L{ROWS + 1}")

    dv_num = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0",
                            allow_blank=True)
    dv_num.error = "Enter a positive number."
    ws.add_data_validation(dv_num)
    dv_num.add(f"D2:D{ROWS + 1}")
    dv_num.add(f"H2:K{ROWS + 1}")

    dv_date = DataValidation(type="date", operator="between",
                             formula1="DATE(2026,1,1)", formula2="DATE(2030,12,31)",
                             allow_blank=True)
    dv_date.error = "Enter a real date (dd/mm/yyyy)."
    ws.add_data_validation(dv_date)
    dv_date.add(f"A2:A{ROWS + 1}")

    rng = f"W2:W{ROWS + 1}"
    ws.conditional_formatting.add(rng, CellIsRule(
        operator="equal", formula=['"Time Exceeded"'], fill=WARN_FILL))
    ws.conditional_formatting.add(rng, CellIsRule(
        operator="equal", formula=['"Clean"'], fill=GOOD_FILL))
    ws.conditional_formatting.add(f"S2:S{ROWS + 1}", CellIsRule(
        operator="equal", formula=['"CHECK TIMES"'], fill=WARN_FILL,
        font=Font(bold=True, color="B42318")))
    ws.conditional_formatting.add(f"O2:O{ROWS + 1}", CellIsRule(
        operator="equal", formula=['"Yes"'], fill=WARN_FILL))
    ws.auto_filter.ref = f"A1:{DECANT_COLS[-1][1]}{ROWS + 1}"
    return ws


# ---------------------------------------------------------------------------
# 3. Shift Plan - planned target vs actual output per shift
# ---------------------------------------------------------------------------
def build_shift_plan(wb: Workbook):
    ws = wb.create_sheet("Shift Plan")
    headers = [
        "Shift Date", "Shift", "Team", "Planned Heads", "Planned Isotainers",
        "Shift Start", "Shift End", "Planned Hours",
        "Actual Heads", "Actual Isotainers", "Actual MT",
        "Variance Isotainers", "Attainment %", "Avg Decant (hrs)",
        "Avg Interval (hrs)", "Total Delay Mins", "Productive Hours",
        "Isotainers per Head", "Shift Key", "Shift Notes",
    ]
    for i, h in enumerate(headers, start=1):
        ws.cell(row=1, column=i, value=h)
    _style_header(ws)
    for i in range(9, 19):
        ws.cell(row=1, column=i).fill = PatternFill("solid", fgColor="334155")

    n = 800
    d = f"'Decant Log'"
    for r in range(2, n + 2):
        key = f'$S{r}'
        ws[f"S{r}"] = f'=IF($A{r}="","",TEXT($A{r},"yyyy-mm-dd")&"|"&$B{r})'
        ws[f"H{r}"] = f'=IF(OR($F{r}="",$G{r}=""),"",ROUND(IF($G{r}<$F{r},$G{r}+1-$F{r},$G{r}-$F{r})*24,2))'
        ws[f"I{r}"] = f'=IF({key}="","",IFERROR(MAXIFS({d}!$D:$D,{d}!$AC:$AC,{key}),""))'
        ws[f"J{r}"] = f'=IF({key}="","",COUNTIFS({d}!$AC:$AC,{key},{d}!$G:$G,"<>"))'
        ws[f"K{r}"] = f'=IF({key}="","",ROUND(SUMIFS({d}!$AA:$AA,{d}!$AC:$AC,{key}),3))'
        ws[f"L{r}"] = f'=IF(OR($E{r}="",$J{r}=""),"",$J{r}-$E{r})'
        ws[f"M{r}"] = f'=IF(OR($E{r}="",$E{r}=0,$J{r}=""),"",ROUND($J{r}/$E{r},4))'
        ws[f"N{r}"] = f'=IF({key}="","",IFERROR(ROUND(AVERAGEIFS({d}!$S:$S,{d}!$AC:$AC,{key},{d}!$S:$S,">0"),2),""))'
        ws[f"O{r}"] = f'=IF({key}="","",IFERROR(ROUND(AVERAGEIFS({d}!$T:$T,{d}!$AC:$AC,{key},{d}!$T:$T,">0"),2),""))'
        ws[f"P{r}"] = f'=IF({key}="","",SUMIFS({d}!$K:$K,{d}!$AC:$AC,{key}))'
        ws[f"Q{r}"] = f'=IF({key}="","",ROUND(SUMIFS({d}!$AB:$AB,{d}!$AC:$AC,{key}),2))'
        ws[f"R{r}"] = f'=IF(OR($I{r}="",$I{r}=0,$J{r}=""),"",ROUND($J{r}/$I{r},3))'

        for c in range(1, len(headers) + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = BORDER
            cell.font = Font(size=10)
            if 9 <= c <= 19:
                cell.fill = LOCK_FILL
        ws[f"A{r}"].number_format = "dd/mm/yyyy"
        ws[f"F{r}"].number_format = "hh:mm"
        ws[f"G{r}"].number_format = "hh:mm"
        ws[f"M{r}"].number_format = "0.0%"
        for col in ("H", "N", "O", "Q", "R"):
            ws[f"{col}{r}"].number_format = "0.00"
        ws[f"K{r}"].number_format = "#,##0.000"

    for col, dv in {"B": _list_dv(SHIFTS, "Shift"), "C": _list_dv(TEAMS, "Team")}.items():
        ws.add_data_validation(dv)
        dv.add(f"{col}2:{col}{n + 1}")

    ws.conditional_formatting.add(f"M2:M{n + 1}", CellIsRule(
        operator="lessThan", formula=["1"], fill=WARN_FILL))
    ws.conditional_formatting.add(f"M2:M{n + 1}", CellIsRule(
        operator="greaterThanOrEqual", formula=["1"], fill=GOOD_FILL))

    _widths(ws, {"A": 12, "B": 12, "C": 8, "D": 12, "E": 15, "F": 11, "G": 11, "H": 12,
                 "I": 11, "J": 13, "K": 12, "L": 13, "M": 12, "N": 13, "O": 14,
                 "P": 13, "Q": 13, "R": 14, "S": 20, "T": 34})
    ws.auto_filter.ref = f"A1:T{n + 1}"
    return ws


# ---------------------------------------------------------------------------
# 4. Delay Log - shift level stoppages not tied to one decant
# ---------------------------------------------------------------------------
def build_delay_log(wb: Workbook):
    ws = wb.create_sheet("Delay Log")
    headers = ["Shift Date", "Shift", "Team", "Area", "Start Date & Time",
               "End Date & Time", "Delay Minutes", "Delay Hours", "Delay Category",
               "Delay Owner", "Isotainer Affected", "Description", "Action Taken",
               "Escalate?", "Escalated To", "Resolved?", "Shift Key"]
    for i, h in enumerate(headers, start=1):
        ws.cell(row=1, column=i, value=h)
    _style_header(ws)
    for i in (7, 8, 17):
        ws.cell(row=1, column=i).fill = PatternFill("solid", fgColor="334155")

    n = 1500
    for r in range(2, n + 2):
        ws[f"G{r}"] = f'=IF(OR($E{r}="",$F{r}=""),"",ROUND(($F{r}-$E{r})*1440,0))'
        ws[f"H{r}"] = f'=IF($G{r}="","",ROUND($G{r}/60,2))'
        ws[f"Q{r}"] = f'=IF($A{r}="","",TEXT($A{r},"yyyy-mm-dd")&"|"&$B{r})'
        for c in range(1, len(headers) + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = BORDER
            cell.font = Font(size=10)
            if c in (7, 8, 17):
                cell.fill = LOCK_FILL
        ws[f"A{r}"].number_format = "dd/mm/yyyy"
        ws[f"E{r}"].number_format = "dd/mm/yyyy hh:mm"
        ws[f"F{r}"].number_format = "dd/mm/yyyy hh:mm"

    areas = ["Decant Station 1", "Decant Station 2", "Warehouse 6", "Yard", "Weighbridge",
             "In Transit", "Safripol Site", "Other"]
    for col, dv in {
        "B": _list_dv(SHIFTS, "Shift"),
        "C": _list_dv(TEAMS, "Team"),
        "D": _list_dv(areas, "Where the delay occurred"),
        "J": _list_dv(DELAY_OWNERS, "Accountable party"),
        "K": _list_dv(ISOS + ["N/A"], "Isotainer affected"),
        "N": _list_dv(YESNO, "Escalate?"),
        "P": _list_dv(YESNO, "Resolved?"),
    }.items():
        ws.add_data_validation(dv)
        dv.add(f"{col}2:{col}{n + 1}")

    dv_cat = DataValidation(type="list",
                            formula1=f"=Lookups!$E$4:$E${3 + len(DELAY_CATEGORIES)}",
                            allow_blank=True, showDropDown=False)
    ws.add_data_validation(dv_cat)
    dv_cat.add(f"I2:I{n + 1}")

    ws.conditional_formatting.add(f"N2:N{n + 1}", CellIsRule(
        operator="equal", formula=['"Yes"'], fill=WARN_FILL))
    _widths(ws, {"A": 12, "B": 12, "C": 8, "D": 18, "E": 18, "F": 18, "G": 12, "H": 11,
                 "I": 26, "J": 16, "K": 16, "L": 40, "M": 34, "N": 10, "O": 18,
                 "P": 11, "Q": 20})
    ws.auto_filter.ref = f"A1:Q{n + 1}"
    return ws


# ---------------------------------------------------------------------------
# 5. Drawdown Plan - the Safripol offtake commitment we measure against
# ---------------------------------------------------------------------------
def build_drawdown_plan(wb: Workbook):
    ws = wb.create_sheet("Drawdown Plan")
    headers = ["Date", "Planned Isotainers", "Planned MT", "Cumulative Planned MT",
               "Cumulative Planned %", "Notes"]
    for i, h in enumerate(headers, start=1):
        ws.cell(row=1, column=i, value=h)
    _style_header(ws)
    for i in (4, 5):
        ws.cell(row=1, column=i).fill = PatternFill("solid", fgColor="334155")

    n = 400
    for r in range(2, n + 2):
        ws[f"C{r}"] = f'=IF($B{r}="","",ROUND($B{r}*Target_Load_KG/1000,2))'
        ws[f"D{r}"] = f'=IF($A{r}="","",ROUND(SUM($C$2:$C{r}),2))'
        ws[f"E{r}"] = f'=IF($D{r}="","",MIN(1,$D{r}/Target_Total_MT))'
        for c in range(1, 7):
            cell = ws.cell(row=r, column=c)
            cell.border = BORDER
            cell.font = Font(size=10)
            if c in (3, 4, 5):
                cell.fill = LOCK_FILL
        ws[f"A{r}"].number_format = "dd/mm/yyyy"
        ws[f"C{r}"].number_format = "#,##0.00"
        ws[f"D{r}"].number_format = "#,##0.00"
        ws[f"E{r}"].number_format = "0.0%"

    _widths(ws, {"A": 13, "B": 18, "C": 14, "D": 20, "E": 19, "F": 44})
    return ws


# ---------------------------------------------------------------------------
# 6. Staff - resourcing linked to shift output
# ---------------------------------------------------------------------------
def build_staff(wb: Workbook, staff_rows: list[tuple]):
    ws = wb.create_sheet("Staff")
    headers = ["Name", "Surname", "Company", "Position", "Department", "Team",
               "Hours Per Day", "Allocation", "Trained Y/N", "Worked Previous Shipment",
               "Active Y/N"]
    for i, h in enumerate(headers, start=1):
        ws.cell(row=1, column=i, value=h)
    _style_header(ws)
    r = 2
    for row in staff_rows:
        for i, v in enumerate(row[:10], start=1):
            cell = ws.cell(row=r, column=i, value=v)
            cell.border = BORDER
            cell.font = Font(size=10)
        ws.cell(row=r, column=11, value="Yes").border = BORDER
        r += 1
    dv = _list_dv(YESNO, "Currently on the shipment?")
    ws.add_data_validation(dv)
    dv.add(f"K2:K{max(r, 200)}")
    _widths(ws, {"A": 16, "B": 16, "C": 12, "D": 24, "E": 14, "F": 8, "G": 13,
                 "H": 60, "I": 12, "J": 22, "K": 11})
    ws.auto_filter.ref = f"A1:K{r - 1}"
    return ws


# ---------------------------------------------------------------------------
# 7. Readme / capture rules
# ---------------------------------------------------------------------------
def build_readme(wb: Workbook):
    ws = wb.create_sheet("How To Use", 0)
    lines = [
        ("SAF OPS TRACKING - " + SHIPMENT, "title"),
        ("PTA outbound drawdown capture workbook. Feeds the daily Safripol Power BI / live HTML report.", "sub"),
        ("", ""),
        ("WHAT CHANGED FROM THE PREVIOUS TRACKER", "h"),
        ("1. Start and End are now full DATE + TIME. The old tracker stored time only, so every nightshift", "b"),
        ("   decant that ran past midnight produced a wrong or negative duration.", "b"),
        ("2. Every duration, variance and performance flag is now a FORMULA. Nothing derivable is typed.", "b"),
        ("3. Delay time is captured as MINUTES with a controlled CATEGORY and OWNER, so time lost can be", "b"),
        ("   summed and split by reason. The old 'Time Loss/Saved' column held text and could not be totalled.", "b"),
        ("4. Heads on Shift is captured per row, and Shift Plan compares planned vs actual isotainers.", "b"),
        ("5. Date, Shift, Team, Isotainer and Delay fields are drop-down validated so the report never has", "b"),
        ("   to clean text. The old sheet mixed real dates with text dates and had trailing spaces.", "b"),
        ("6. Drawdown Plan holds the Safripol commitment so % complete and offtake pace are measurable.", "b"),
        ("", ""),
        ("WHAT YOU DO NOT CAPTURE HERE", "h"),
        ("Transit time, Safripol dwell and full turnaround are NOT typed in. They are derived automatically", "b"),
        ("from the GPS tracking feed (Safripol v Connect Dwells) and the Safripol Stock Report. Capturing", "b"),
        ("them by hand would duplicate data and introduce error.", "b"),
        ("", ""),
        ("SHEETS", "h"),
        ("Decant Log     One row per isotainer decant. This is the main capture sheet.", "b"),
        ("Shift Plan     One row per date + shift. Enter planned heads and planned isotainers before the", "b"),
        ("               shift. Actuals pull through automatically from the Decant Log.", "b"),
        ("Delay Log      Shift level stoppages that are not tied to a single decant.", "b"),
        ("Drawdown Plan  Safripol daily offtake commitment. Drives % complete and pace vs plan.", "b"),
        ("Staff          Team allocation. Links resourcing to shift output.", "b"),
        ("Lookups        Targets and drop-down lists. Change a target here and the whole book re-rates.", "b"),
        ("", ""),
        ("CAPTURE RULES", "h"),
        ("- Shift Date is the OPERATIONAL date the shift started on, even if the decant finished after midnight.", "b"),
        ("- Enter Start and End as dd/mm/yyyy hh:mm. If End is before Start the sheet shows CHECK TIMES in red.", "b"),
        ("- Delay Minutes is the time genuinely lost inside that decant. Leave blank or 0 if there was none.", "b"),
        ("- If a stoppage affects the whole shift rather than one isotainer, log it in Delay Log instead.", "b"),
        ("- Set Escalate? to Yes for anything management must action. These surface on the report exception panel.", "b"),
        ("- Grey columns are calculated. Do not type in them.", "b"),
        ("", ""),
        ("TARGETS (edit on the Lookups sheet)", "h"),
        (f"- Allowed decant time: {ALLOWED_DECANT_HRS} hrs      - Target interval between decants: {ALLOWED_INTERVAL_HRS} hrs", "b"),
        (f"- Nominal payload: {TARGET_LOAD_KG:,} kg           - Site dwell target: {SITE_TARGET_HRS} hrs", "b"),
        (f"- Transit target: {TRANSIT_TARGET_HRS} hrs                - Shipment target: {TARGET_TOTAL_MT:,.2f} MT", "b"),
    ]
    r = 1
    for text, kind in lines:
        c = ws.cell(row=r, column=1, value=text)
        if kind == "title":
            c.font = Font(bold=True, size=18, color=ACCENT)
            ws.row_dimensions[r].height = 26
        elif kind == "sub":
            c.font = Font(italic=True, size=11, color="475467")
        elif kind == "h":
            c.font = Font(bold=True, size=11, color="FFFFFF")
            c.fill = HEAD_FILL
        else:
            c.font = Font(size=10, color="1D2939")
        r += 1
    ws.column_dimensions["A"].width = 118
    ws.sheet_view.showGridLines = False
    return ws


def load_staff(path: str) -> list[tuple]:
    if not path or not os.path.exists(path):
        return []
    import shutil
    import tempfile
    from openpyxl import load_workbook
    tmp = os.path.join(tempfile.gettempdir(), "_staff_src.xlsx")
    shutil.copy2(path, tmp)
    wb = load_workbook(tmp, data_only=True)
    ws = wb[wb.sheetnames[0]]
    out = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[0]:
            continue
        if str(row[0]).strip().lower() == "name":     # repeated header block
            continue
        out.append(tuple((str(v).replace("\xa0", " ").strip() if isinstance(v, str) else v)
                         for v in row))
    return out


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", required=True)
    ap.add_argument("--staff", default="")
    args = ap.parse_args()

    wb = Workbook()
    wb.remove(wb.active)

    build_decant_log(wb)
    build_shift_plan(wb)
    build_delay_log(wb)
    build_drawdown_plan(wb)
    build_staff(wb, load_staff(args.staff))
    build_lookups(wb)
    build_readme(wb)
    wb.move_sheet("How To Use", offset=-10)

    wb.properties.title = f"SAF Ops Tracking - {SHIPMENT}"
    wb.properties.creator = "Connect Logistics - Process Optimization & Development"
    os.makedirs(os.path.dirname(os.path.abspath(args.out)), exist_ok=True)
    wb.save(args.out)
    print(f"Written: {args.out}")


if __name__ == "__main__":
    main()
